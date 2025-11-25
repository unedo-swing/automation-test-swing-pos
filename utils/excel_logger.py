import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
import os

class ExcelLogger:
    def __init__(self, file_path="reports/test_result.xlsx"):
        self.file_path = file_path
        self.steps = []
        self.status = "PASS"
        self.error_message = ""
        self.screenshot_path = ""

        # Ensure folder exists
        os.makedirs(os.path.dirname(file_path), exist_ok=True)

        # Create file if not exist
        if not os.path.exists(file_path):
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.title = "Test Results"

            headers = ["Date", "Test Name", "Steps", "Status", "Error Message", "Screenshot"]
            sheet.append(headers)

            # Style header
            for col in range(1, len(headers) + 1):
                cell = sheet.cell(1, col)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4F81BD", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            wb.save(file_path)

    def add_step(self, description):
        self.steps.append(description)

    def mark_failed(self, message, screenshot_path):
        self.status = "FAIL"
        self.error_message = message
        self.screenshot_path = screenshot_path

    def save(self, test_name):
        wb = openpyxl.load_workbook(self.file_path)
        sheet = wb.active

        # Number the steps
        numbered_steps = [f"{i}. {step}" for i, step in enumerate(self.steps, start=1)]

        row = [
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            test_name,
            "\n".join(numbered_steps),
            self.status,
            self.error_message,
            self.screenshot_path if self.screenshot_path else "-"
        ]

        sheet.append(row)

        # Auto width
        for column_cells in sheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = length + 4

        wb.save(self.file_path)
