from openpyxl import load_workbook
import pandas as pd

def read_excel_data(file_path, sheet_name):
    import openpyxl
    wb = openpyxl.load_workbook(file_path)
    sheet = wb[sheet_name]
    data = []

    headers = [cell.value for cell in sheet[1]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(row):  # skip empty rows
            data.append(dict(zip(headers, row)))
    return data


def write_excel_result(file_path, sheet_name, row_num, result, steps="", message=""):
    wb = load_workbook(file_path)
    sheet = wb[sheet_name]

    headers = [cell.value for cell in sheet[1]]

    def ensure_column(col_name):
        if col_name in headers:
            return headers.index(col_name) + 1
        else:
            new_col = len(headers) + 1
            sheet.cell(row=1, column=new_col, value=col_name)
            headers.append(col_name)
            return new_col

    result_col = ensure_column("STATUS")
    steps_col = ensure_column("TEST_STEPS")
    message_col = ensure_column("MESSAGE")

    sheet.cell(row=row_num, column=result_col, value=result)
    sheet.cell(row=row_num, column=steps_col, value=steps)
    sheet.cell(row=row_num, column=message_col, value=message)

    wb.save(file_path)


def read_localization_data(file_path, sheet_name="Sheet1"):
    df = pd.read_excel(file_path, sheet_name=sheet_name)
    # Make sure to strip any extra spaces
    df["LOCKEY_ID"] = df["LOCKEY_ID"].astype(str).str.strip()
    df["VALUE"] = df["VALUE"].astype(str).str.strip()
    return dict(zip(df["LOCKEY_ID"], df["VALUE"]))

